from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def build_workbook(title, headers, rows):
    """Первая строка — заголовки на русском; даты уже отформатированы как ДД.ММ.ГГГГ."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
